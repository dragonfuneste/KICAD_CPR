import openpyxl
import time
from Pricing_lcsc import get_lcsc_price_from_page
def Update_Price_Stock(name,nom_feuille):
    wb = openpyxl.load_workbook(name)
    ws = wb[nom_feuille]

    # on repère la colonne de chaque en-tête (ligne 1) une seule fois
    entetes = {}
    for cell in ws[1]:
        if cell.value:
            entetes[str(cell.value).strip()] = cell.column  # numéro de colonne (1-based)

    col_nom = entetes["nom du composant"]
    col_lcsc = entetes["Ref LCSC"]

    # si les colonnes Price / Stock n'existent pas encore, on les crée à la fin
    if "Price" not in entetes:
        entetes["Price"] = ws.max_column + 1
        ws.cell(row=1, column=entetes["Price"], value="Price")
    if "Stock website" not in entetes:
        entetes["Stock website"] = ws.max_column + 1
        ws.cell(row=1, column=entetes["Stock website"], value="Stock website")

    col_price = entetes["Price"]
    col_stock = entetes["Stock website"]

    # on parcourt les lignes de données (à partir de la ligne 2)
    for row_idx in range(2, ws.max_row + 1):
        ref_lcsc_cell = ws.cell(row=row_idx, column=col_lcsc)
        ref_lcsc = ref_lcsc_cell.value

        if ref_lcsc is None or str(ref_lcsc).strip() == "":
            continue

        ref_lcsc = str(ref_lcsc).strip()
        nom = ws.cell(row=row_idx, column=col_nom).value
        print(nom, ref_lcsc)

        try:
            result = get_lcsc_price_from_page(ref_lcsc, 1)
        except Exception as e:
            print(f"Erreur inattendue pour {ref_lcsc} : {e}")
            result = None

        if result is None:
            print(f"  -> aucun résultat pour {ref_lcsc}, on passe à la suite")
            time.sleep(1)
            continue

        # on écrit uniquement la valeur, l'hyperlien de la cellule Ref LCSC
        # n'est jamais touché puisqu'on ne modifie que les colonnes Price/Stock
        ws.cell(row=row_idx, column=col_price, value=result["prix_unitaire_eur"])
        ws.cell(row=row_idx, column=col_stock, value=result["stock"])

        time.sleep(1)

    wb.save(name)


import pandas as pd
from find_component import find_by_row
from component_checker import clean_BOM


def build_pricing_report(bom_path, library_path, n=5, sheet_name_lib="Feuille 1"):
    """
    Calcule le coût total du BOM pour fabriquer n cartes, vérifie la
    disponibilité en stock, et détaille le poids de chaque famille de
    composant (colonne 'type' de la bibliothèque) dans le coût total.

    Retourne (report, cout_total, composants_manquants, repartition_par_famille)
    """
    BOM = clean_BOM(bom_path)
    library = pd.read_excel(library_path, sheet_name=sheet_name_lib)

    lignes = []
    cout_total = 0.0

    for index, row in BOM.iterrows():
        resultat = find_by_row(library, row)

        qte_par_carte = row.get("Quantity Per PCB", 0)
        try:
            qte_par_carte = float(qte_par_carte)
        except (ValueError, TypeError):
            qte_par_carte = 0

        qte_necessaire = qte_par_carte * n

        composant = None
        if isinstance(resultat, pd.Series):
            composant = resultat
        elif isinstance(resultat, pd.DataFrame) and not resultat.empty:
            composant = resultat.iloc[0]

        ligne_base = {
            "Ligne BOM": index + 1,
            "References": row.get("References", ""),
            "Manufacturer Ref": row.get("Manufacturer Ref", ""),
            "Quantité/carte": qte_par_carte,
            f"Quantité pour {n} cartes": qte_necessaire,
        }

        if composant is None:
            lignes.append({
                **ligne_base,
                "Famille": "inconnue",
                "Composant bibliothèque": "",
                "Prix unitaire (EUR)": None,
                "Coût ligne (EUR)": None,
                "Stock": None,
                "Disponible": "inconnu (pas de match)",
                "Manque": qte_necessaire,
            })
            continue

        prix_unitaire = composant.get("Price")
        stock = composant.get("Stock website")
        famille = composant.get("type")
        famille = famille if pd.notna(famille) else "inconnue"

        cout_ligne = None
        if pd.notna(prix_unitaire):
            cout_ligne = prix_unitaire * qte_necessaire
            cout_total += cout_ligne

        if pd.notna(stock):
            manque = max(0, qte_necessaire - stock)
            statut_dispo = "oui" if manque == 0 else "non"
        else:
            manque = qte_necessaire
            statut_dispo = "inconnu (pas de stock enregistré)"

        lignes.append({
            **ligne_base,
            "Famille": famille,
            "Composant bibliothèque": composant["nom du composant"],
            "Prix unitaire (EUR)": prix_unitaire,
            "Coût ligne (EUR)": cout_ligne,
            "Stock": stock,
            "Disponible": statut_dispo,
            "Manque": manque,
        })

    report = pd.DataFrame(lignes)
    composants_manquants = report[report["Disponible"] != "oui"].copy()

    # --------------------------------------------------------
    # Répartition du coût par famille de composant
    # --------------------------------------------------------
    repartition_par_famille = (
        report.dropna(subset=["Coût ligne (EUR)"])
        .groupby("Famille")["Coût ligne (EUR)"]
        .sum()
        .reset_index()
        .rename(columns={"Coût ligne (EUR)": "Coût total (EUR)"})
        .sort_values("Coût total (EUR)", ascending=False)
    )

    if cout_total > 0:
        repartition_par_famille["% du coût total"] = (
            repartition_par_famille["Coût total (EUR)"] / cout_total * 100
        ).round(1)
    else:
        repartition_par_famille["% du coût total"] = 0.0

    # --------------------------------------------------------
    # Résumé console
    # --------------------------------------------------------
    print(f"Coût total pour {n} carte(s) : {cout_total:.2f} EUR")

    if composants_manquants.empty:
        print("Tous les composants sont disponibles en quantité suffisante.")
    else:
        print(f"{len(composants_manquants)} ligne(s) à problème :")
        print(composants_manquants[
            ["Manufacturer Ref", "Composant bibliothèque", f"Quantité pour {n} cartes", "Stock", "Manque", "Disponible"]
        ].to_string(index=False))

    print("\nRépartition du coût par famille de composant :")
    print(repartition_par_famille.to_string(index=False))

    return report, cout_total, composants_manquants, repartition_par_famille