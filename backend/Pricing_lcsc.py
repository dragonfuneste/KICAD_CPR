import requests
import re
import json
import time
import openpyxl
import time
# ============================================================
# CONVERSION DE DEVISE (USD -> EUR), taux mis en cache
# ============================================================

_taux_cache = {"taux": None, "expire_at": 0}

def get_taux_usd_vers_eur():
    """Récupère le taux de change USD->EUR (source BCE), mis en cache 1h."""

    if _taux_cache["taux"] is not None and time.time() < _taux_cache["expire_at"]:
        return _taux_cache["taux"]

    try:
        response = requests.get(
            "https://api.frankfurter.app/latest",
            params={"from": "USD", "to": "EUR"},
            timeout=10,
        )
        response.raise_for_status()
        data = response.json()
        taux = data["rates"]["EUR"]
    except Exception as e:
        print("ERREUR récupération taux de change :", e)
        return None

    _taux_cache["taux"] = taux
    _taux_cache["expire_at"] = time.time() + 3600
    return taux


def convertir_usd_vers_eur(montant_usd):
    """Convertit un montant en USD vers EUR."""
    if montant_usd is None:
        return None
    taux = get_taux_usd_vers_eur()
    if taux is None:
        return None
    return round(montant_usd * taux, 5)

import requests
import re
import json

HEADERS_NAVIGATEUR = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                  "(KHTML, like Gecko) Chrome/120.0 Safari/537.36",
    "Accept-Language": "fr-FR,fr;q=0.9",
}


# ============================================================
# LCSC (scraping) - prix EUR + stock
# ============================================================

def get_lcsc_price_from_page(product_code, quantite=1):
    """Récupère prix (EUR) et stock en scrapant la page produit LCSC."""

    url = f"https://www.lcsc.com/product-detail/{product_code}.html"

    try:
        response = requests.get(url, headers=HEADERS_NAVIGATEUR, timeout=15)
        print("LCSC - HTTP :", response.status_code)
        response.raise_for_status()
    except requests.exceptions.RequestException as e:
        print("ERREUR réseau LCSC :", e)
        return None

    html = response.text

    prix_match = re.findall(r'"productPriceList"\s*:\s*(\[[^\]]*\])', html)
    stock_match = re.search(r'"stockNumber"\s*:\s*(\d+)', html)

    if not prix_match:
        print("LCSC : bloc de prix introuvable dans le HTML brut (probablement chargé en JS).")
        return None

    try:
        paliers_usd = json.loads(prix_match[0])
    except json.JSONDecodeError:
        print("LCSC : bloc de prix trouvé mais JSON invalide.")
        return None

    prix_paliers_eur = []
    for p in paliers_usd:
        prix_eur = convertir_usd_vers_eur(p.get("usdPrice"))
        if prix_eur is not None:
            prix_paliers_eur.append((p.get("ladder"), prix_eur))
    prix_paliers_eur.sort(key=lambda x: x[0])

    prix = prix_paliers_eur[0][1] if prix_paliers_eur else None
    for q, p in prix_paliers_eur:
        if q <= quantite:
            prix = p
        else:
            break

    stock = int(stock_match.group(1)) if stock_match else None

    return {
        "reference_lcsc": product_code,
        "quantite_demandee": quantite,
        "prix_unitaire_eur": prix,
        "prix_paliers_eur": prix_paliers_eur,
        "stock": stock,
        "quantite_suffisante": (stock is not None and stock >= quantite),
    }


def Update_Price_Stock(name,nom_feuille = "Feuille 1"):
    wb = openpyxl.load_workbook(name)
    ws = wb[nom_feuille]

    # on repère la colonne de chaque en-tête (ligne 1) une seule fois
    entetes = {}
    for cell in ws[1]:
        if cell.value:
            entetes[str(cell.value).strip()] = cell.column  # numéro de colonne (1-based)

    col_nom = entetes["Manufacturer Ref"]
    col_lcsc = entetes["reference_LCSC"]

    # si les colonnes Price / Stock n'existent pas encore, on les crée à la fin
    if "Price" not in entetes:
        entetes["Price"] = ws.max_column + 1
        ws.cell(row=1, column=entetes["Price"], value="Price")
    if "Stock website" not in entetes:
        entetes["Stock website"] = ws.max_column + 1
        ws.cell(row=1, column=entetes["Stock website"], value="Stock website")

    col_price = entetes["Price"]
    col_stock = entetes["Stock website"]

    # on parcourt les lignes de données (à partir de la ligne 2)
    for row_idx in range(2, ws.max_row + 1):
        ref_lcsc_cell = ws.cell(row=row_idx, column=col_lcsc)
        ref_lcsc = ref_lcsc_cell.value

        if ref_lcsc is None or str(ref_lcsc).strip() == "":
            continue

        ref_lcsc = str(ref_lcsc).strip()
        nom = ws.cell(row=row_idx, column=col_nom).value
        print(nom, ref_lcsc)

        try:
            result = get_lcsc_price_from_page(ref_lcsc, 1)
        except Exception as e:
            print(f"Erreur inattendue pour {ref_lcsc} : {e}")
            result = None

        if result is None:
            print(f"  -> aucun résultat pour {ref_lcsc}, on passe à la suite")
            time.sleep(1)
            continue

        # on écrit uniquement la valeur, l'hyperlien de la cellule Ref LCSC
        # n'est jamais touché puisqu'on ne modifie que les colonnes Price/Stock
        ws.cell(row=row_idx, column=col_price, value=result["prix_unitaire_eur"])
        ws.cell(row=row_idx, column=col_stock, value=result["stock"])

        time.sleep(1)

    wb.save(name)

